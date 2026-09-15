"""Modelo de escenarios inmobiliarios. Genera contexto/inmuebles/escenarios-deptos.xlsx
con supuestos editables (celdas amarillas) y fórmulas. Los valores impresos por consola
son los mismos que se citan en analisis-tres-deptos.md."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "contexto" / "inmuebles" / "escenarios-deptos.xlsx"

# ---------- supuestos (fuente y fecha en el comentario) ----------
A = dict(
    tc=1500,                 # BNA 07-09-2026: 1.480 compra / 1.530 venta. Referencia, no afirmación.
    uva=2106.13,             # BCRA 07-09-2026
    depto_a=145_000,         # dicho por Agustina: donde viven, 65 m2
    depto_b=140_000,         # dicho por Agustina: en pago, mismas condiciones
    saldo_b=0,               # DESCONOCIDO: cuánto falta pagar de B y en qué moneda/tasa
    depto_c=140_000,         # objetivo tercer depto, "mismas condiciones"
    m2_grande=95,            # 3-4 amb para mudarse
    usd_m2_grande=2500,      # Caballito abr-2026: 1.800 (sur) a 2.800 (Primera Junta / P. Rivadavia)
    costo_compra=0.08,       # sellos 1,75% (no es vivienda única) + escribano ~2% + comisión 4%+IVA ~ 8-8,6%
    costo_venta=0.051,       # comisión 3%+IVA 3,63% + ITI 1,5%
    alquiler_2amb=900_000,   # $/mes 2 amb 65 m2 Caballito; Zonaprop abr-2026 promedio 2 amb $834k CABA, Caballito 600-800k (35-45 m2)
    vacancia=1/12,           # un mes por año
    gastos_prop=0.10,        # ABL, mantenimiento, gestión, sobre el alquiler bruto
    tna_1ra=0.115,           # BBVA sep-2026 1ra vivienda: UVA + 11,5%
    tna_2da=0.15,            # Provincia / BBVA 2da vivienda o no cobra sueldo: UVA + 15%
    plazo_meses=360,
    ltv=0.75,
    cuota_ingreso=0.25,      # bancos: 20-30%
    ort=426_000,             # recibo de sueldo, único ingreso demostrable hoy
    roxom_usdt=1600,         # NO demostrable hasta que se facture por monotributo
    roxom_pedido=2400,
    cftea_personal=0.90,     # préstamo personal sep-2026: 58,9% (premium) a 320%; Macro/Bancor ~91-93%
)

def cuota_por_millon(tna, n):
    r = tna / 12
    return 1e6 * r / (1 - (1 + r) ** -n)

def capacidad(ingreso_pesos, tna, n=A["plazo_meses"], ratio=A["cuota_ingreso"], tc=A["tc"]):
    cuota_max = ingreso_pesos * ratio
    prestamo = cuota_max / cuota_por_millon(tna, n) * 1e6
    return cuota_max, prestamo / tc

# ---------- cálculos que se citan ----------
res = {}
alq_neto_mes = A["alquiler_2amb"] * (1 - A["vacancia"]) * (1 - A["gastos_prop"])
res["alq_bruto_usd"] = A["alquiler_2amb"] / A["tc"]
res["alq_neto_usd"] = alq_neto_mes / A["tc"]
res["yield_bruto"] = A["alquiler_2amb"] * 12 / A["tc"] / A["depto_b"]
res["yield_neto"] = alq_neto_mes * 12 / A["tc"] / A["depto_b"]

# Escenario 1: comprar C con hipoteca
costo_c = A["depto_c"] * (1 + A["costo_compra"])
prestamo_c = A["depto_c"] * A["ltv"]
efectivo_c = costo_c - prestamo_c
for tag, tna in (("1ra", A["tna_1ra"]), ("2da", A["tna_2da"])):
    cuota = prestamo_c * A["tc"] / 1e6 * cuota_por_millon(tna, A["plazo_meses"])
    res[f"cuota_c_{tag}_ars"] = cuota
    res[f"cuota_c_{tag}_usd"] = cuota / A["tc"]
    res[f"flujo_c_{tag}_usd"] = res["alq_neto_usd"] - cuota / A["tc"]
    res[f"cobertura_{tag}"] = alq_neto_mes / cuota
res["costo_c"] = costo_c; res["prestamo_c"] = prestamo_c; res["efectivo_c"] = efectivo_c
res["ingreso_min_c_1ra"] = res["cuota_c_1ra_ars"] / A["cuota_ingreso"]
res["ingreso_min_c_2da"] = res["cuota_c_2da_ars"] / A["cuota_ingreso"]

# Escenario 2: vender A y B, comprar uno grande
grande = A["m2_grande"] * A["usd_m2_grande"]
neto_venta = (A["depto_a"] + A["depto_b"]) * (1 - A["costo_venta"]) - A["saldo_b"]
costo_grande = grande * (1 + A["costo_compra"])
res["grande"] = grande; res["neto_venta"] = neto_venta; res["costo_grande"] = costo_grande
res["friccion_s2"] = (A["depto_a"] + A["depto_b"]) * A["costo_venta"] + grande * A["costo_compra"]
res["sobra_s2"] = neto_venta - costo_grande

# Escenario 3: quedarse con A y B, alquilar B, no comprar
res["flujo_s3_usd"] = res["alq_neto_usd"]
res["anios_ahorro_grande"] = None  # se calcula abajo con ahorro mensual

# Capacidad de crédito de Agustina
esc_ing = {
    "Hoy: sólo ORT ($426k demostrable)": A["ort"],
    "ORT + Roxom 1.600 facturado (6-12 m de facturas)": A["ort"] + A["roxom_usdt"] * A["tc"],
    "ORT + Roxom 2.400 facturado": A["ort"] + A["roxom_pedido"] * A["tc"],
}
cap = {}
for k, ing in esc_ing.items():
    for tag, tna in (("1ra", A["tna_1ra"]), ("2da", A["tna_2da"])):
        cap[(k, tag)] = capacidad(ing, tna)

# Préstamo personal: costo mensual vs renta
tem_personal = (1 + A["cftea_personal"]) ** (1/12) - 1
res["tem_personal"] = tem_personal
res["tem_alquiler"] = res["yield_neto"] / 12

# ---------- salida consola ----------
print(f"Alquiler 2 amb: bruto US$ {res['alq_bruto_usd']:.0f}/mes, neto US$ {res['alq_neto_usd']:.0f}/mes")
print(f"Yield sobre US$ {A['depto_b']:,}: bruto {res['yield_bruto']:.2%}, neto {res['yield_neto']:.2%}")
print(f"\nE1 comprar C: costo total US$ {costo_c:,.0f}, préstamo {A['ltv']:.0%} US$ {prestamo_c:,.0f}, efectivo US$ {efectivo_c:,.0f}")
for tag in ("1ra", "2da"):
    print(f"  UVA+{A['tna_'+tag]:.1%}: cuota inicial ${res['cuota_c_'+tag+'_ars']:,.0f} = US$ {res['cuota_c_'+tag+'_usd']:.0f}; "
          f"alquiler cubre {res['cobertura_'+tag]:.0%}; flujo US$ {res['flujo_c_'+tag+'_usd']:+.0f}/mes; "
          f"ingreso neto demostrable mínimo ${res['ingreso_min_c_'+tag]:,.0f}")
print(f"\nE2 vender A+B y comprar {A['m2_grande']} m2 a {A['usd_m2_grande']}: precio US$ {grande:,.0f}, con costos {costo_grande:,.0f}")
print(f"  neto de vender A+B: US$ {neto_venta:,.0f} (saldo B pendiente = {A['saldo_b']}: DESCONOCIDO)")
print(f"  fricción total (comisiones, sellos, ITI, escribano): US$ {res['friccion_s2']:,.0f}; sobra/falta: US$ {res['sobra_s2']:+,.0f}")
print(f"\nE3 alquilar B, no comprar: US$ {res['flujo_s3_usd']:.0f}/mes neto")
print("\nCapacidad hipotecaria de Agustina (cuota <= 25% del neto demostrable, 30 años):")
for (k, tag), (cm, pu) in cap.items():
    print(f"  {k:52s} UVA+{A['tna_'+tag]:.1%}: cuota máx ${cm:,.0f} -> préstamo US$ {pu:,.0f}")
print(f"\nPréstamo personal CFTEA {A['cftea_personal']:.0%} = {tem_personal:.2%} mensual vs renta neta {res['tem_alquiler']:.2%} mensual: "
      f"carry {res['tem_alquiler']-tem_personal:+.2%}/mes")
print(f"Cuota por millón: 1ra {cuota_por_millon(A['tna_1ra'],360):,.0f}, 2da {cuota_por_millon(A['tna_2da'],360):,.0f}")

# ---------- xlsx con fórmulas ----------
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Supuestos"
Y = PatternFill("solid", fgColor="FFF2CC"); B = Font(bold=True)
ws.append(["Supuestos — celdas amarillas editables. Todo lo demás se recalcula."]); ws["A1"].font = B
ws.append([])
rows = [
    ("tc", "Tipo de cambio de referencia", A["tc"], "BNA 07-09-2026: 1.480/1.530. No hay TC verificado en resúmenes."),
    ("uva", "Valor UVA", A["uva"], "BCRA 07-09-2026"),
    ("depto_a", "Depto A (viven) US$", A["depto_a"], "Dicho por Agustina. 65 m2, Plaza Irlanda."),
    ("depto_b", "Depto B (en pago) US$", A["depto_b"], "Dicho por Agustina. Cid Campeador."),
    ("saldo_b", "Saldo pendiente de B (US$)", A["saldo_b"], "DESCONOCIDO. Pedir a Maxi: cuánto falta, moneda, tasa, plazo."),
    ("depto_c", "Depto C objetivo US$", A["depto_c"], "Mismas condiciones que B."),
    ("m2_grande", "m2 del depto grande", A["m2_grande"], "3-4 ambientes"),
    ("usd_m2_grande", "US$/m2 Caballito", A["usd_m2_grande"], "Abr-2026: 1.800 sur a 2.800 Primera Junta"),
    ("costo_compra", "Costos de compra (% del precio)", A["costo_compra"], "Sellos 1,75% + escribano ~2% + comisión 4%+IVA"),
    ("costo_venta", "Costos de venta (% del precio)", A["costo_venta"], "Comisión 3%+IVA + ITI 1,5%"),
    ("alquiler", "Alquiler 2 amb 65 m2 ($/mes)", A["alquiler_2amb"], "Zonaprop abr-2026: 2 amb CABA $834k; Caballito 600-800k para 35-45 m2. Verificar con relevamiento."),
    ("vacancia", "Vacancia (fracción del año)", A["vacancia"], "1 mes por año"),
    ("gastos_prop", "Gastos del propietario (% del alquiler)", A["gastos_prop"], "ABL, mantenimiento, gestión"),
    ("tna_1ra", "TNA hipotecario 1ra vivienda (+UVA)", A["tna_1ra"], "BBVA sep-2026"),
    ("tna_2da", "TNA hipotecario 2da vivienda / inversión (+UVA)", A["tna_2da"], "Provincia sin cuenta sueldo / 2da vivienda"),
    ("plazo", "Plazo (meses)", A["plazo_meses"], ""),
    ("ltv", "Financiación máxima (LTV)", A["ltv"], "75-80% según banco"),
    ("ratio", "Cuota / ingreso neto máximo", A["cuota_ingreso"], "Bancos: 20-30%"),
    ("ort", "ORT neto ($/mes, demostrable)", A["ort"], "Recibo de sueldo"),
    ("roxom", "Roxom (USDT/mes)", A["roxom_usdt"], "NO demostrable hasta facturar por monotributo"),
    ("roxom_pedido", "Roxom si sale la review (USDT/mes)", A["roxom_pedido"], ""),
    ("cftea", "CFTEA préstamo personal", A["cftea_personal"], "Sep-2026: 58,9% premium a 320%"),
]
names = {}
for key, label, val, src in rows:
    ws.append([label, val, src]); r = ws.max_row
    ws.cell(r, 2).fill = Y; names[key] = f"Supuestos!$B${r}"
    if isinstance(val, float) and val < 1: ws.cell(r, 2).number_format = "0.00%"
    else: ws.cell(r, 2).number_format = "#,##0.00"
ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 80
N = names

def pmt(tna, plazo, capital_ars):
    return f"-PMT({tna}/12,{plazo},{capital_ars})"

# Hoja escenarios
es = wb.create_sheet("Escenarios")
es.append(["Escenarios comparados", "", "", "Cómo se calcula"]); es["A1"].font = B
es.append([])
es.append(["Alquiler 2 amb", "", "", ""]); es.cell(es.max_row,1).font = B
es.append(["Alquiler bruto US$/mes", f"={N['alquiler']}/{N['tc']}", "", "alquiler / TC"])
es.append(["Alquiler neto US$/mes", f"={N['alquiler']}*(1-{N['vacancia']})*(1-{N['gastos_prop']})/{N['tc']}", "", "neto de vacancia y gastos"]); r_neto = es.max_row
es.append(["Yield bruto anual", f"=B{r_neto-1}*12/{N['depto_b']}", "", "sobre el valor del depto"])
es.append(["Yield neto anual", f"=B{r_neto}*12/{N['depto_b']}", "", ""]); r_yn = es.max_row
es.append([])
es.append(["E1 — Comprar C con hipoteca (viven en A, alquilan B y C)"]); es.cell(es.max_row,1).font = B
es.append(["Costo total de C con gastos US$", f"={N['depto_c']}*(1+{N['costo_compra']})"]); r_cc = es.max_row
es.append(["Préstamo (LTV) US$", f"={N['depto_c']}*{N['ltv']}"]); r_pc = es.max_row
es.append(["Efectivo necesario US$", f"=B{r_cc}-B{r_pc}", "", "precio + gastos − préstamo. ¿De dónde sale?"])
for tag in ("1ra", "2da"):
    es.append([f"Cuota inicial UVA+{tag} $", "=" + pmt(N['tna_'+tag], N['plazo'], f"B{r_pc}*{N['tc']}"), "", "sube con UVA cada mes"]); rc = es.max_row
    es.append([f"Cuota inicial UVA+{tag} US$", f"=B{rc}/{N['tc']}"])
    es.append([f"Alquiler neto cubre la cuota ({tag})", f"=B{r_neto}*{N['tc']}/B{rc}"]); es.cell(es.max_row,2).number_format = "0%"
    es.append([f"Flujo mensual con C alquilado ({tag}) US$", f"=B{r_neto}-B{rc}/{N['tc']}", "", "negativo = hay que poner plata todos los meses"])
    es.append([f"Ingreso neto demostrable mínimo del tomador ({tag}) $", f"=B{rc}/{N['ratio']}", "", "cuota / 25%"])
es.append(["Flujo total E1: renta B + renta C − cuota C (2da vivienda) US$", f"=B{r_neto}*2-B{rc}/{N['tc']}", "", "lo que queda para la pareja cada mes"])
es.append([])
es.append(["E2 — Vender A y B, comprar uno grande"]); es.cell(es.max_row,1).font = B
es.append(["Precio depto grande US$", f"={N['m2_grande']}*{N['usd_m2_grande']}"]); r_g = es.max_row
es.append(["Costo con gastos de compra US$", f"=B{r_g}*(1+{N['costo_compra']})"]); r_gc = es.max_row
es.append(["Neto de vender A y B US$", f"=({N['depto_a']}+{N['depto_b']})*(1-{N['costo_venta']})-{N['saldo_b']}"]); r_nv = es.max_row
es.append(["Fricción total (se pierde) US$", f"=({N['depto_a']}+{N['depto_b']})*{N['costo_venta']}+B{r_g}*{N['costo_compra']}", "", "comisiones, sellos, ITI, escribano"])
es.append(["Sobra (+) / falta (−) US$", f"=B{r_nv}-B{r_gc}"])
es.append(["Renta que se pierde al vender B US$/mes", f"=B{r_neto}"])
es.append([])
es.append(["E3 — Quedarse con A y B, alquilar B, no comprar"]); es.cell(es.max_row,1).font = B
es.append(["Flujo mensual US$", f"=B{r_neto}"])
es.append(["Ahorro mensual adicional de la pareja US$ (editar)", 500]); es.cell(es.max_row,2).fill = Y; r_ah = es.max_row
es.append(["Meses para juntar el efectivo de C", f"=B{r_cc}-B{r_pc}", "", "efectivo necesario"]); r_ef = es.max_row
es.cell(r_ef,1).value = "Efectivo necesario para C (75% LTV) US$"
es.append(["Meses para juntarlo con renta B + ahorro", f"=B{r_ef}/(B{r_neto}+B{r_ah})"])
es.column_dimensions["A"].width = 58; es.column_dimensions["B"].width = 16; es.column_dimensions["D"].width = 50
for row in es.iter_rows(min_row=3, min_col=2, max_col=2):
    c = row[0]
    if c.number_format == "General": c.number_format = "#,##0"
es.cell(r_yn-1,2).number_format = "0.00%"; es.cell(r_yn,2).number_format = "0.00%"

# Hoja capacidad de crédito
cp = wb.create_sheet("Capacidad Agustina")
cp.append(["Capacidad hipotecaria de Agustina — cuota ≤ ratio × ingreso neto demostrable"]); cp["A1"].font = B
cp.append(["Escenario de ingreso", "Ingreso neto demostrable $", "Cuota máx $", f"Préstamo US$ (1ra)", f"Préstamo US$ (2da)", "Condición"])
for c in cp[2]: c.font = B
data = [
    ("Hoy: sólo ORT", f"={N['ort']}", "En mora en Galicia Mastercard: ningún banco precalifica hasta salir y que BCRA lo refleje."),
    ("ORT + Roxom 1.600 facturado", f"={N['ort']}+{N['roxom']}*{N['tc']}", "Requiere 6-12 meses de facturas de monotributo. Cada banco aplica su descuento."),
    ("ORT + Roxom 2.400 facturado", f"={N['ort']}+{N['roxom_pedido']}*{N['tc']}", "Si sale la review y se factura."),
]
for label, ing, cond in data:
    cp.append([label, ing, None, None, None, cond]); r = cp.max_row
    cp.cell(r,3).value = f"=B{r}*{N['ratio']}"
    cp.cell(r,4).value = f"=PV({N['tna_1ra']}/12,{N['plazo']},-C{r})/{N['tc']}"
    cp.cell(r,5).value = f"=PV({N['tna_2da']}/12,{N['plazo']},-C{r})/{N['tc']}"
    for col in (2,3,4,5): cp.cell(r,col).number_format = "#,##0"
cp.append([])
cp.append(["Préstamo personal en vez de hipoteca"]); cp.cell(cp.max_row,1).font = B
cp.append(["Costo mensual del préstamo (TEM desde CFTEA)", f"=(1+{N['cftea']})^(1/12)-1"]); cp.cell(cp.max_row,2).number_format = "0.00%"; r_tem = cp.max_row
cp.append(["Renta neta mensual del depto", f"=Escenarios!B{r_yn}/12"]); cp.cell(cp.max_row,2).number_format = "0.00%"
cp.append(["Carry mensual (renta − costo)", f"=B{r_tem+1}-B{r_tem}"]); cp.cell(cp.max_row,2).number_format = "0.00%"
cp.column_dimensions["A"].width = 48; cp.column_dimensions["F"].width = 90
for col in "BCDE": cp.column_dimensions[col].width = 20

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("\nEscrito", OUT)
